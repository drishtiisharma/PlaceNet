from pathlib import Path
from io import BytesIO

import fitz
from docx import Document

from app.models.parsed_resume import ParsedResume


class ParserService:

    @staticmethod
    def parse(
        file_bytes: bytes,
        filename: str
    ) -> ParsedResume:

        extension = Path(filename).suffix.lower()

        if extension == ".pdf":
            text = ParserService._parse_pdf(file_bytes)

        elif extension == ".docx":
            text = ParserService._parse_docx(file_bytes)

        elif extension == ".txt":
            text = file_bytes.decode("utf-8")

        else:
            raise ValueError(
                f"Unsupported file type: {extension}"
            )

        text = text.strip()

        if not text:
            raise ValueError(
                "No readable text found in resume."
            )

        candidate_name = ParserService._extract_candidate_name(text)

        return ParsedResume(
            candidate_name=candidate_name,
            resume_text=text
        )

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:

        pdf = fitz.open(
            stream=file_bytes,
            filetype="pdf"
        )

        pages = []

        for page in pdf:
            pages.append(page.get_text())

        pdf.close()

        return "\n".join(pages)

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> str:

        document = Document(BytesIO(file_bytes))

        text = []

        # Paragraphs
        for paragraph in document.paragraphs:

            value = paragraph.text.strip()

            if value:
                text.append(value)

        # Tables
        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    value = cell.text.strip()

                    if value:
                        text.append(value)

        return "\n".join(text)

    @staticmethod
    def _extract_candidate_name(text: str) -> str:

        for line in text.splitlines():

            line = line.strip()

            if line:
                return line.replace("/", "_")

        return "Unknown_Candidate"